"""
Resume Parser Service for RecruitFlow AI (Phase 4).
Extracts raw text and metadata from PDF and DOCX files using PyMuPDF and python-docx.
Includes scanned/image-only PDF detection, encrypted document handling,
table content extraction, and clean text normalization.
"""
import re
from pathlib import Path
from typing import Dict, Any, List

try:
    import pymupdf as fitz
except ImportError:
    import fitz

import docx


class ResumeParserError(Exception):
    """Custom exception for resume parsing failures."""
    def __init__(self, message, error_code="parsing_error"):
        super().__init__(message)
        self.error_code = error_code


class ResumeParserService:
    """Extracts raw text and structural metadata from PDF and DOCX documents."""

    @classmethod
    def parse_document(cls, file_path: str, file_type: str = None) -> Dict[str, Any]:
        """
        Parses document based on extension/type.
        Returns dictionary with:
        - raw_text: cleaned string
        - page_count: int
        - char_count: int
        - word_count: int
        - warnings: list of string warnings
        """
        path = Path(file_path)
        if not path.exists():
            raise ResumeParserError(f"File not found: {file_path}", "file_not_found")

        ext = (file_type or path.suffix).lower().lstrip(".")

        if ext == "pdf":
            return cls.parse_pdf(str(path))
        elif ext == "docx":
            return cls.parse_docx(str(path))
        else:
            raise ResumeParserError(f"Unsupported document format: {ext}", "unsupported_format")

    @classmethod
    def parse_pdf(cls, file_path: str) -> Dict[str, Any]:
        """Extracts text from PDF document using PyMuPDF with page and scanned-image diagnostics."""
        warnings: List[str] = []
        pages_text: List[str] = []
        page_count = 0

        try:
            doc = fitz.open(file_path)
        except Exception as e:
            raise ResumeParserError(f"Failed to open PDF document: {str(e)}", "corrupt_pdf")

        try:
            if doc.is_encrypted:
                raise ResumeParserError("The PDF document is password-protected or encrypted.", "encrypted_pdf")

            page_count = doc.page_count
            if page_count == 0:
                warnings.append("Document has 0 pages.")
                return {
                    "raw_text": "",
                    "page_count": 0,
                    "char_count": 0,
                    "word_count": 0,
                    "warnings": warnings,
                }

            for page_num in range(page_count):
                page = doc.load_page(page_num)
                text = page.get_text("text") or ""
                pages_text.append(text)

        finally:
            doc.close()

        full_text = "\n\n".join(pages_text)
        cleaned_text = cls.clean_text(full_text)
        char_count = len(cleaned_text)
        word_count = len(cleaned_text.split())

        # Check for scanned or image-based PDF (very low text density)
        if char_count < 50:
            warnings.append("Low text content detected. Document may be a scanned image or photograph without OCR text.")

        return {
            "raw_text": cleaned_text,
            "page_count": page_count,
            "char_count": char_count,
            "word_count": word_count,
            "warnings": warnings,
        }

    @classmethod
    def parse_docx(cls, file_path: str) -> Dict[str, Any]:
        """Extracts text from DOCX document including body paragraphs and table cells."""
        warnings: List[str] = []
        text_parts: List[str] = []

        try:
            doc = docx.Document(file_path)
        except Exception as e:
            raise ResumeParserError(f"Failed to open DOCX document: {str(e)}", "corrupt_docx")

        # Paragraphs
        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if text:
                text_parts.append(text)

        # Tables (resumes frequently format experience/skills in tables)
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                # Deduplicate identical adjacent cells (from merged cells)
                deduped = []
                for c in row_cells:
                    if not deduped or c != deduped[-1]:
                        deduped.append(c)
                if deduped:
                    text_parts.append(" | ".join(deduped))

        full_text = "\n".join(text_parts)
        cleaned_text = cls.clean_text(full_text)
        char_count = len(cleaned_text)
        word_count = len(cleaned_text.split())

        if char_count < 50:
            warnings.append("Low text content detected in DOCX document.")

        # Estimate page count for DOCX (roughly 500 words per page)
        estimated_pages = max(1, (word_count // 500) + (1 if word_count % 500 > 0 else 0)) if word_count > 0 else 1

        return {
            "raw_text": cleaned_text,
            "page_count": estimated_pages,
            "char_count": char_count,
            "word_count": word_count,
            "warnings": warnings,
        }

    @staticmethod
    def clean_text(text: str) -> str:
        """
        Normalizes unicode characters, replaces non-standard whitespace,
        bullet points, and fixes broken linebreaks.
        """
        if not text:
            return ""

        # Normalize line endings
        text = text.replace("\r\n", "\n").replace("\r", "\n")

        # Replace non-breaking spaces and special spaces
        text = text.replace("\xa0", " ").replace("\u200b", "")

        # Standardize bullet symbols
        bullet_symbols = ["\u2022", "\u2023", "\u25e6", "\u2043", "\u2219", "\u25cf", "\u25cb", "\u25aa", "\u25ab"]
        for b in bullet_symbols:
            text = text.replace(b, "- ")

        # Remove excessive whitespace within lines
        lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.split("\n")]

        # Collapse 3+ consecutive newlines to 2
        cleaned = "\n".join(lines)
        cleaned = re.sub(r"\n{3,}", "\n\n", cleaned)

        return cleaned.strip()
