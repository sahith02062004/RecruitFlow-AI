"""
Automated Test Suite for Phase 4: Resume Processing and Extraction.
Tests:
- Magic byte validation for PDF and DOCX
- File extension and size constraints
- Deterministic extraction: skills taxonomy, contact info, education, experience
- PyMuPDF and python-docx text parsers
- Candidate upload and automated profile enrichment
- Candidate correction and review workflow
- Role-based authorization: tenant-isolated recruiter access, candidate ownership, forbidden downloads
- Authorized document download and safe deletion
"""
import io
import os
import tempfile
from pathlib import Path
import pytest
from app.core.extensions import db
from app.models.models import Resume, Job, Application, Organization, OrganizationMembership, CandidateProfile
from app.services.resume_storage_service import ResumeStorageService, ResumeStorageError
from app.services.resume_parser_service import ResumeParserService
from app.services.resume_extraction_service import ResumeExtractionService
from app.services.resume_processing_service import ResumeProcessingService

try:
    import pymupdf as fitz
except ImportError:
    import fitz

import docx


def create_sample_pdf_bytes(
    name="Jordan Smith",
    email="jordan.smith@example.com",
    phone="+1 (555) 987-6543",
    skills="Python, Flask, PostgreSQL, Docker, AWS, Kubernetes",
    body_text=None
) -> bytes:
    """Helper creating a valid PDF document with text in memory."""
    doc = fitz.open()
    page = doc.new_page()
    content = body_text or f"""
{name}
{email} | {phone}
https://www.linkedin.com/in/jordansmith
https://github.com/jordansmith

Summary
Passionate senior software engineer with 5+ years of experience building scalable backend platforms.

Skills
{skills}

Experience
Senior Backend Engineer - CloudTech Inc
2020 - Present
- Architected resilient microservices in Python and Flask.
- Orchestrated container workloads on AWS with Kubernetes and Docker.

Software Engineer - InnovateSoft
2018 - 2020
- Built PostgreSQL database pipelines and REST APIs.

Education
B.Tech in Computer Science
2014 - 2018
"""
    page.insert_text((50, 72), content)
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


def create_sample_docx_bytes(
    name="Taylor Swift",
    email="taylor.swift@example.com",
    phone="555-888-9999",
    skills="React, TypeScript, Node.js, GraphQL, Redis"
) -> bytes:
    """Helper creating a valid DOCX document in memory."""
    doc = docx.Document()
    doc.add_heading(name, level=1)
    doc.add_paragraph(f"{email} | {phone}")
    doc.add_paragraph("https://www.linkedin.com/in/taylorswift")
    doc.add_paragraph("https://github.com/taylorswift")

    doc.add_heading("Professional Summary", level=2)
    doc.add_paragraph("Experienced full-stack web developer with 4 years of experience delivering modern web applications.")

    doc.add_heading("Technical Skills", level=2)
    doc.add_paragraph(skills)

    doc.add_heading("Work History", level=2)
    doc.add_paragraph("Full Stack Engineer | 2020 - 2024")
    doc.add_paragraph("- Developed dynamic web applications with React, TypeScript, and Node.js.")

    doc.add_heading("Education", level=2)
    table = doc.add_table(rows=1, cols=2)
    row_cells = table.rows[0].cells
    row_cells[0].text = "B.S. in Software Engineering"
    row_cells[1].text = "2016 - 2020"

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# -------------------------------------------------------------
# Storage & Magic Byte Validation Tests
# -------------------------------------------------------------

def test_resume_storage_magic_byte_validation_pdf(app):
    """Verify valid PDF magic header is accepted and spoofed PDF is rejected."""
    with app.app_context():
        valid_pdf_bytes = create_sample_pdf_bytes()
        valid_file = io.BytesIO(valid_pdf_bytes)
        valid_file.filename = "valid_resume.pdf"

        meta = ResumeStorageService.validate_file(valid_file)
        assert meta["extension"] == ".pdf"
        assert meta["mime_type"] == "application/pdf"

        # Spoofed file: plain text disguised as .pdf
        fake_file = io.BytesIO(b"This is just plain text, not a real PDF file!")
        fake_file.filename = "spoofed_resume.pdf"
        with pytest.raises(ResumeStorageError) as exc:
            ResumeStorageService.validate_file(fake_file)
        assert exc.value.error_code == "invalid_pdf_signature"


def test_resume_storage_magic_byte_validation_docx(app):
    """Verify valid DOCX is accepted and non-ZIP file disguised as .docx is rejected."""
    with app.app_context():
        valid_docx_bytes = create_sample_docx_bytes()
        valid_file = io.BytesIO(valid_docx_bytes)
        valid_file.filename = "valid_resume.docx"

        meta = ResumeStorageService.validate_file(valid_file)
        assert meta["extension"] == ".docx"
        assert "wordprocessingml" in meta["mime_type"]

        # Spoofed file: plain text disguised as .docx
        fake_file = io.BytesIO(b"Not a zip file or docx archive!")
        fake_file.filename = "spoofed_resume.docx"
        with pytest.raises(ResumeStorageError) as exc:
            ResumeStorageService.validate_file(fake_file)
        assert exc.value.error_code == "invalid_docx_signature"


def test_resume_storage_rejects_unsupported_extensions(app):
    """Verify executable, script, or image formats are strictly rejected."""
    with app.app_context():
        bad_file = io.BytesIO(b"malicious script content")
        bad_file.filename = "payload.exe"
        with pytest.raises(ResumeStorageError) as exc:
            ResumeStorageService.validate_file(bad_file)
        assert exc.value.error_code == "invalid_extension"


def test_resume_storage_rejects_oversized_file(app):
    """Verify file exceeding 5MB limit is rejected."""
    with app.app_context():
        huge_bytes = b"%PDF-1.4\n" + (b"0" * (6 * 1024 * 1024))
        huge_file = io.BytesIO(huge_bytes)
        huge_file.filename = "oversized_resume.pdf"
        with pytest.raises(ResumeStorageError) as exc:
            ResumeStorageService.validate_file(huge_file)
        assert exc.value.error_code == "file_too_large"


# -------------------------------------------------------------
# Parser & Diagnostics Tests
# -------------------------------------------------------------

def test_resume_parser_service_pdf():
    """Test text parsing from PDF using PyMuPDF."""
    with tempfile.TemporaryDirectory() as td:
        pdf_bytes = create_sample_pdf_bytes(name="Samantha Fox", skills="Python, Django, React")
        file_path = Path(td) / "test_sample.pdf"
        file_path.write_bytes(pdf_bytes)

        result = ResumeParserService.parse_pdf(str(file_path))
        assert result["page_count"] >= 1
        assert "Samantha Fox" in result["raw_text"]
        assert "Python" in result["raw_text"]
        assert result["char_count"] > 100
        assert len(result["warnings"]) == 0


def test_resume_parser_service_docx():
    """Test text parsing from DOCX using python-docx."""
    with tempfile.TemporaryDirectory() as td:
        docx_bytes = create_sample_docx_bytes(name="Marcus Chen", skills="FastAPI, MySQL, Redis")
        file_path = Path(td) / "test_sample.docx"
        file_path.write_bytes(docx_bytes)

        result = ResumeParserService.parse_docx(str(file_path))
        assert "Marcus Chen" in result["raw_text"]
        assert "FastAPI" in result["raw_text"]
        assert "B.S. in Software Engineering" in result["raw_text"]
        assert result["char_count"] > 100


def test_resume_parser_scanned_pdf_warning():
    """Test diagnostic warning when PDF has virtually no text (scanned image)."""
    with tempfile.TemporaryDirectory() as td:
        doc = fitz.open()
        doc.new_page()  # Blank page with 0 text
        pdf_bytes = doc.tobytes()
        doc.close()

        file_path = Path(td) / "blank.pdf"
        file_path.write_bytes(pdf_bytes)

        result = ResumeParserService.parse_pdf(str(file_path))
        assert any("Low text content" in w for w in result["warnings"])


# -------------------------------------------------------------
# Deterministic Extraction Service Tests
# -------------------------------------------------------------

def test_deterministic_extraction_service():
    """Verify name, contact, skills taxonomy normalization, and experience extraction."""
    sample_text = """
Alex Morgan
alex.morgan@testdomain.com | +1 (555) 345-6789
linkedin.com/in/alexmorgan | github.com/alexmorgan

Professional Summary
Senior engineer with 6+ years of experience leading teams in building high throughput backend APIs.

Technical Skills
Proficient in Python, py, JS, TypeScript, postgres, k8s, Docker, Flask, AWS, PyTorch, and REST APIs.

Work Experience
Senior Software Architect
2019 - Present
- Designed microservices architecture.

Software Developer
2017 - 2019
- Built web applications and database systems.

Education
Bachelor of Technology in Computer Science
2013 - 2017
"""
    extracted = ResumeExtractionService.extract_all(sample_text, fallback_name="Default Name")

    assert extracted["name"] == "Alex Morgan"
    assert extracted["email"] == "alex.morgan@testdomain.com"
    assert "555" in extracted["phone"]
    assert extracted["links"].get("linkedin") == "https://www.linkedin.com/in/alexmorgan"
    assert extracted["links"].get("github") == "https://github.com/alexmorgan"

    # Verify skill taxonomy normalization (aliases mapped to canonical names)
    skills = extracted["skills"]
    assert "Python" in skills
    assert "JavaScript" in skills  # Normalized from "JS"
    assert "TypeScript" in skills
    assert "PostgreSQL" in skills  # Normalized from "postgres"
    assert "Kubernetes" in skills  # Normalized from "k8s"
    assert "Docker" in skills
    assert "Flask" in skills
    assert "AWS" in skills
    assert "PyTorch" in skills
    assert "REST API" in skills

    # Experience calculation
    assert extracted["total_experience_years"] >= 6.0

    # Education degree extraction
    edu_degrees = [e["degree"] for e in extracted["education"]]
    assert any("B.Tech" in d or "Bachelor" in d for d in edu_degrees)


# -------------------------------------------------------------
# Candidate Upload, Ingestion & Correction Route Tests
# -------------------------------------------------------------

def test_candidate_uploads_resume_and_syncs_profile(client, candidate_user):
    """Test candidate successfully uploads PDF resume with auto-sync to profile."""
    # Login as candidate
    client.post("/auth/login", data={"email": candidate_user.email, "password": "CandidateSecret123!"})

    pdf_bytes = create_sample_pdf_bytes(
        name="Alex Candidate",
        email=candidate_user.email,
        phone="+1 555-432-1098",
        skills="Python, Flask, Docker, Kubernetes, MySQL, React"
    )

    data = {
        "resume_file": (io.BytesIO(pdf_bytes), "alex_candidate_cv.pdf"),
        "sync_profile": "y"
    }

    response = client.post("/resumes/upload", data=data, content_type="multipart/form-data", follow_redirects=True)
    assert response.status_code == 200
    assert b"uploaded and parsed successfully" in response.data

    # Check resume in DB
    resume = Resume.query.filter_by(candidate_user_id=candidate_user.id).first()
    assert resume is not None
    assert resume.processing_status == "processed"
    assert resume.original_filename == "alex_candidate_cv.pdf"
    assert "Python" in resume.skills_list
    assert "Docker" in resume.skills_list

    # Check candidate profile was synchronized
    profile = candidate_user.candidate_profile
    assert profile.phone is not None
    assert "Docker" in profile.skills


def test_candidate_corrects_extracted_data(client, candidate_user):
    """Test candidate reviews and submits corrections to parsed resume data."""
    client.post("/auth/login", data={"email": candidate_user.email, "password": "CandidateSecret123!"})

    # Create initial resume record
    pdf_bytes = create_sample_pdf_bytes(name="Alex Candidate")
    data = {
        "resume_file": (io.BytesIO(pdf_bytes), "resume_to_correct.pdf"),
        "sync_profile": "y"
    }
    client.post("/resumes/upload", data=data, content_type="multipart/form-data")
    resume = Resume.query.filter_by(candidate_user_id=candidate_user.id).first()
    assert resume is not None

    # Submit correction
    corr_data = {
        "name": "Alexander Candidate Jr.",
        "email": candidate_user.email,
        "phone": "+1 555-999-0000",
        "skills": "Python, Flask, PyTorch, LangChain, MySQL, Docker",
        "total_experience_years": "5.5",
        "linkedin_url": "https://www.linkedin.com/in/alexandercandidate",
        "summary": "Updated senior AI systems builder summary.",
        "sync_profile": "y"
    }
    res_corr = client.post(f"/resumes/{resume.id}/correct", data=corr_data, follow_redirects=True)
    assert res_corr.status_code == 200
    assert b"Resume details updated successfully" in res_corr.data

    # Check updated extracted data
    db.session.refresh(resume)
    data_dict = resume.get_extracted_data()
    assert data_dict["name"] == "Alexander Candidate Jr."
    assert "LangChain" in resume.skills_list
    assert resume.parsed_experience_years == 5.5


# -------------------------------------------------------------
# Security & Access Control Tests
# -------------------------------------------------------------

def test_resume_access_control_isolation(client, candidate_user, recruiter_user, session):
    """
    Test access control:
    - Another user cannot access candidate's resume (HTTP 403).
    - Recruiter cannot view candidate's resume until candidate applies to recruiter's job.
    - After application, recruiter gains authorized access.
    """
    # 1. Candidate uploads resume
    client.post("/auth/login", data={"email": candidate_user.email, "password": "CandidateSecret123!"})
    pdf_bytes = create_sample_pdf_bytes()
    client.post("/resumes/upload", data={"resume_file": (io.BytesIO(pdf_bytes), "private_cv.pdf")}, content_type="multipart/form-data")
    resume = Resume.query.filter_by(candidate_user_id=candidate_user.id).first()
    assert resume is not None

    # Log out candidate
    client.get("/auth/logout")

    # 2. Recruiter logs in before application: access forbidden
    client.post("/auth/login", data={"email": recruiter_user.email, "password": "RecruiterSecret123!"})
    res = client.get(f"/resumes/{resume.id}")
    assert res.status_code == 403

    # Recruiter cannot download either
    res_dl = client.get(f"/resumes/{resume.id}/download")
    assert res_dl.status_code == 403

    # 3. Create job in recruiter's organization and candidate applies
    recruiter_org = recruiter_user.memberships.first().organization
    job = Job(
        organization_id=recruiter_org.id,
        created_by_user_id=recruiter_user.id,
        title="AI Engineer",
        department="Engineering",
        workplace_type="remote",
        employment_type="full-time",
        minimum_experience_years=2.0,
        required_skills="Python, PyTorch",
        description="Build state-of-the-art AI systems.",
        salary_currency="USD",
        status="published"
    )
    session.add(job)
    session.flush()

    app_record = Application(
        job_id=job.id,
        candidate_user_id=candidate_user.id,
        resume_id=resume.id,
        status="submitted"
    )
    session.add(app_record)
    session.commit()

    # 4. Now recruiter views candidate's resume detail and downloads file
    res_detail = client.get(f"/resumes/{resume.id}")
    assert res_detail.status_code == 200
    assert b"private_cv.pdf" in res_detail.data

    res_dl_authorized = client.get(f"/resumes/{resume.id}/download")
    assert res_dl_authorized.status_code == 200
    assert "attachment" in res_dl_authorized.headers.get("Content-Disposition", "")
    res_dl_authorized.close()


def test_candidate_downloads_and_deletes_own_resume(client, candidate_user):
    """Test candidate can download and delete their own resume document."""
    client.post("/auth/login", data={"email": candidate_user.email, "password": "CandidateSecret123!"})

    pdf_bytes = create_sample_pdf_bytes()
    client.post("/resumes/upload", data={"resume_file": (io.BytesIO(pdf_bytes), "cv_to_delete.pdf")}, content_type="multipart/form-data")
    resume = Resume.query.filter_by(candidate_user_id=candidate_user.id).first()
    assert resume is not None
    resume_id = resume.id

    # Download
    dl_res = client.get(f"/resumes/{resume_id}/download")
    assert dl_res.status_code == 200
    dl_res.close()

    # Delete
    del_res = client.post(f"/resumes/{resume_id}/delete", follow_redirects=True)
    assert del_res.status_code == 200
    assert b"deleted successfully" in del_res.data

    # Verify gone from database
    assert db.session.get(Resume, resume_id) is None
