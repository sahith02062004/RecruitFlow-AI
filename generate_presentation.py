import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_deck(output_path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    NAVY = RGBColor(15, 23, 42)       # #0f172a
    BLUE = RGBColor(30, 58, 138)      # #1e3a8a
    ACCENT_TEAL = RGBColor(15, 118, 110) # #0f766e
    BG_LIGHT = RGBColor(248, 250, 252) # #f8fafc
    WHITE = RGBColor(255, 255, 255)
    GRAY = RGBColor(100, 116, 139)    # #64748b
    CARD_BG = RGBColor(255, 255, 255)
    BORDER_COLOR = RGBColor(226, 232, 240) # #e2e8f0
    GREEN = RGBColor(22, 101, 52)     # #166534

    blank_layout = prs.slide_layouts[6]

    def set_slide_background(slide, color):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = color

    def add_header(slide, title_text, category_text="RECRUITFLOW AI — MASTER ARCHITECTURE"):
        # Category / breadcrumb
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.5), Inches(0.4))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = ACCENT_TEAL
        p_cat.font.name = "Arial"

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.5), Inches(0.8))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(26)
        p_title.font.bold = True
        p_title.font.color.rgb = NAVY
        p_title.font.name = "Arial"

        # Horizontal accent bar
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.45), Inches(11.7), Inches(0.04))
        bar.fill.solid()
        bar.fill.fore_color.rgb = BORDER_COLOR
        bar.line.fill.background()

    def add_card(slide, left, top, width, height, title, items, badge=""):
        # Card shape
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = BORDER_COLOR
        card.line.width = Pt(1.5)

        # Card text
        tb = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.2), width - Inches(0.4), height - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True

        p0 = tf.paragraphs[0]
        p0.text = title
        p0.font.size = Pt(16)
        p0.font.bold = True
        p0.font.color.rgb = BLUE
        p0.font.name = "Arial"
        p0.space_after = Pt(8)

        if badge:
            # Badge text below title
            p_badge = tf.add_paragraph()
            p_badge.text = f"★ {badge}"
            p_badge.font.size = Pt(11)
            p_badge.font.bold = True
            p_badge.font.color.rgb = ACCENT_TEAL
            p_badge.space_after = Pt(10)

        for item in items:
            p = tf.add_paragraph()
            p.text = f"•  {item}"
            p.font.size = Pt(12)
            p.font.color.rgb = NAVY
            p.font.name = "Arial"
            p.space_after = Pt(6)

    # ==========================================================
    # SLIDE 1: TITLE SLIDE (Dark Navy Premium Theme)
    # ==========================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1, NAVY)

    # Accent decorative bar
    top_bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.2), Inches(1.5), Inches(2.5), Inches(0.1))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = ACCENT_TEAL
    top_bar.line.fill.background()

    # Title box
    tbox = s1.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(11.0), Inches(2.2))
    tf1 = tbox.text_frame
    tf1.word_wrap = True
    p1 = tf1.paragraphs[0]
    p1.text = "RecruitFlow AI"
    p1.font.size = Pt(44)
    p1.font.bold = True
    p1.font.color.rgb = WHITE
    p1.font.name = "Arial"

    p2 = tf1.add_paragraph()
    p2.text = "AI-Assisted Resume Screening, Ranking & Candidate Intelligence Platform"
    p2.font.size = Pt(22)
    p2.font.color.rgb = RGBColor(147, 197, 253) # light blue
    p2.font.name = "Arial"
    p2.space_before = Pt(8)

    # Subtitle / Details Card
    sub_box = s1.shapes.add_textbox(Inches(1.2), Inches(4.3), Inches(11.0), Inches(2.2))
    tf_sub = sub_box.text_frame
    tf_sub.word_wrap = True
    
    details = [
        ("Architecture:", "Python 3.12 Flask Full-Stack + MySQL 8 + SQLAlchemy 2.0"),
        ("AI / ML Engine:", "100% Deterministic Explainable Scoring (Zero external paid APIs, Zero Hallucinations)"),
        ("Validation & Quality:", "85/85 Automated Pytest Test Suite Passing (100% Pass Rate in ~27s)"),
        ("Complete Phases:", "All 7 Phases Implemented End-to-End (Auth, Jobs, Resumes, Matching, Pipeline, Analytics)"),
        ("Presented by:", "Full Stack & Generative AI Software Engineering Portfolio")
    ]
    for label, val in details:
        p = tf_sub.add_paragraph()
        run1 = p.add_run()
        run1.text = f"{label:22} "
        run1.font.bold = True
        run1.font.color.rgb = RGBColor(226, 232, 240)
        run1.font.size = Pt(14)
        run2 = p.add_run()
        run2.text = val
        run2.font.color.rgb = RGBColor(203, 213, 225)
        run2.font.size = Pt(14)
        p.space_after = Pt(4)

    # ==========================================================
    # SLIDE 2: THE PROBLEM & WHY RECRUITFLOW AI
    # ==========================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2, BG_LIGHT)
    add_header(s2, "Problem Statement & Product Vision", "MARKET CHALLENGE & SOLUTION")

    add_card(s2, Inches(0.8), Inches(1.8), Inches(3.6), Inches(5.0),
             "1. Current ATS Friction",
             [
                 "Black-box algorithmic rejection without explaining reasons to candidates.",
                 "Naïve keyword counters fail on synonyms or punish candidates unfairly.",
                 "Disguised file uploads (malicious scripts) create security risks.",
                 "High cost of external commercial LLM APIs with non-deterministic scoring."
             ], "Industry Flaws")

    add_card(s2, Inches(4.8), Inches(1.8), Inches(3.6), Inches(5.0),
             "2. The RecruitFlow AI Vision",
             [
                 "100% Explainable & Transparent candidate evaluation.",
                 "Dual-score separation: ATS readability vs. Job suitability.",
                 "Dignified feedback: Evidence-based 'Why Was My Resume Not Shortlisted?'.",
                 "Ethical AI: Suggestions never hallucinate unearned skills or credentials."
             ], "The Solution")

    add_card(s2, Inches(8.8), Inches(1.8), Inches(3.7), Inches(5.0),
             "3. Business & Technical Impact",
             [
                 "Zero external API dependency saves cloud budget & guarantees 0ms latency.",
                 "Strict tenant isolation keeps corporate client data segregated.",
                 "Human-in-the-loop guarantee: AI never rejects or hires autonomously.",
                 "OWASP hardened: CSRF protection, magic-byte checks & CSV sanitization."
             ], "Key Advantages")

    # ==========================================================
    # SLIDE 3: SYSTEM ARCHITECTURE & TECH STACK
    # ==========================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3, BG_LIGHT)
    add_header(s3, "Enterprise Full-Stack Architecture", "SYSTEM ARCHITECTURE")

    add_card(s3, Inches(0.8), Inches(1.8), Inches(3.6), Inches(5.0),
             "Backend & Database Layer",
             [
                 "Python 3.12+ with Flask Application Factory pattern.",
                 "SQLAlchemy 2.0 ORM with modular relational schema.",
                 "MySQL 8 primary production database (UTF8mb4).",
                 "SQLite in-memory test isolation for sub-30s CI/CD.",
                 "Flask-Migrate / Alembic for reliable versioned migrations."
             ], "Core Engine")

    add_card(s3, Inches(4.8), Inches(1.8), Inches(3.6), Inches(5.0),
             "Security & Tenancy Hardening",
             [
                 "Flask-Login with strong session protection.",
                 "Flask-WTF CSRF protection on 100% of mutation forms.",
                 "Werkzeug scrypt password hashing.",
                 "Multi-tenant OrganizationMembership RBAC decorators.",
                 "Single-use SHA-256 hashed password reset tokens."
             ], "Enterprise Security")

    add_card(s3, Inches(8.8), Inches(1.8), Inches(3.7), Inches(5.0),
             "Frontend & Presentation Layer",
             [
                 "Responsive Bootstrap 5 UI + Bootstrap Icons.",
                 "Enterprise design system: Slate Navy & Professional Teal.",
                 "Chart.js interactive visualizations for hiring funnels.",
                 "Truthful zero-state macros: no fake statistics.",
                 "Interactive Kanban board for live application progression."
             ], "UI / UX System")

    # ==========================================================
    # SLIDE 4: THE 7 CORE MODULES & WORKFLOW
    # ==========================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4, BG_LIGHT)
    add_header(s4, "Complete 7-Phase Milestone Execution", "FEATURE ROADMAP & SCOPE")

    add_card(s4, Inches(0.8), Inches(1.8), Inches(2.7), Inches(5.0),
             "Phases 1 & 2: Core & Auth",
             [
                 "App factory & modular Blueprints.",
                 "Role-Based Access Control (Candidate, Recruiter, Admin).",
                 "Multi-tenant company accounts.",
                 "Secure signup, login, password reset."
             ], "Phase 1 & 2")

    add_card(s4, Inches(3.8), Inches(1.8), Inches(2.7), Inches(5.0),
             "Phases 3 & 4: Jobs & Ingestion",
             [
                 "Job requisition lifecycle (Draft, Published, Closed).",
                 "Candidate application submission & duplicate prevention.",
                 "Magic-byte inspection (PDF / DOCX).",
                 "PyMuPDF + python-docx extraction.",
                 "Canonical taxonomy: 80+ skills."
             ], "Phase 3 & 4")

    add_card(s4, Inches(6.8), Inches(1.8), Inches(2.7), Inches(5.0),
             "Phases 5 & 6: AI & Pipeline",
             [
                 "Deterministic Explainable Scoring.",
                 "Dynamic weight redistribution.",
                 "SHA-256 staleness tracking.",
                 "Multi-round interview scheduler.",
                 "1-5 interviewer evaluation rubrics.",
                 "Confidential internal feedback."
             ], "Phase 5 & 6")

    add_card(s4, Inches(9.8), Inches(1.8), Inches(2.7), Inches(5.0),
             "Phase 7: ATS & Analytics",
             [
                 "7-Category ATS Quality Scorer.",
                 "Dual-score separation architecture.",
                 "10-criteria comparison engine.",
                 "3-tier prioritized writing checklist.",
                 "Transparent shortlisting gap diagnosis.",
                 "Funnel metrics & sanitized CSV export."
             ], "Phase 7")

    # ==========================================================
    # SLIDE 5: RESUME INGESTION & PARSING PIPELINE
    # ==========================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5, BG_LIGHT)
    add_header(s5, "High-Performance Document Ingestion Pipeline", "RESUME PROCESSING")

    add_card(s5, Inches(0.8), Inches(1.8), Inches(3.6), Inches(5.0),
             "1. Upload & Threat Defense",
             [
                 "Strict 5 MB server-side file size enforcement.",
                 "Allowed extensions restricted to .pdf and .docx.",
                 "Magic-byte validation: %PDF- and PK\\x03\\x04 OpenXML.",
                 "Rejection of disguised executables & malicious scripts.",
                 "UUIDv4 random storage keys outside web root to stop path traversal."
             ], "Security Barrier")

    add_card(s5, Inches(4.8), Inches(1.8), Inches(3.6), Inches(5.0),
             "2. Dual-Engine Text Extraction",
             [
                 "PyMuPDF (pymupdf): High-speed PDF text parsing & page counts.",
                 "Scanned PDF warning: Detects image-only files with <50 chars.",
                 "python-docx: Full Word parsing across paragraphs & table cells.",
                 "Deterministic Regex & Pattern Matching: Name, RFC email, phone, links, degrees, experience."
             ], "Parsing Engine")

    add_card(s5, Inches(8.8), Inches(1.8), Inches(3.7), Inches(5.0),
             "3. Canonical Taxonomy & Sync",
             [
                 "80+ Canonical Skills Taxonomy with alias resolution.",
                 "Normalizes: 'k8s' -> Kubernetes, 'postgres' -> PostgreSQL, 'js' -> JavaScript.",
                 "One-click Candidate Profile Synchronization.",
                 "Candidate review & correction modal before submission."
             ], "Normalization")

    # ==========================================================
    # SLIDE 6: EXPLAINABLE AI SCORING ENGINE
    # ==========================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6, BG_LIGHT)
    add_header(s6, "Deterministic & Explainable AI Matching", "SCORING ALGORITHMS")

    add_card(s6, Inches(0.8), Inches(1.8), Inches(3.6), Inches(5.0),
             "Configurable Category Weights",
             [
                 "Required Skills Alignment: 50%",
                 "Preferred Skills Alignment: 20%",
                 "Experience Years Alignment: 20%",
                 "Educational Qualification: 10%",
                 "Strict Total: 100% (1.00) mathematically verified."
             ], "Core Weightings")

    add_card(s6, Inches(4.8), Inches(1.8), Inches(3.6), Inches(5.0),
             "Dynamic Weight Redistribution",
             [
                 "No penalty for omitted job criteria.",
                 "If a job specifies no preferred skills, that 20% is proportionally redistributed across active categories.",
                 "Ensures candidates are never penalized for unrequested attributes.",
                 "Pure Python deterministic engine — zero LLM hallucination or drift."
             ], "Fairness Engine")

    add_card(s6, Inches(8.8), Inches(1.8), Inches(3.7), Inches(5.0),
             "Strict Token Boundary Matching",
             [
                 "Regex word boundaries prevent substring confusion:",
                 "• 'Java' != 'JavaScript'",
                 "• 'C' != 'C++' or 'C#'",
                 "Evidence Coverage Tracking: If evidence is missing, reports 'Unavailable' instead of 0% or 100%.",
                 "SHA-256 Input Hash flags stale scores automatically."
             ], "Audit Safeguards")

    # ==========================================================
    # SLIDE 7: DUAL-SCORE SEPARATION & ATS QUALITY ANALYZER
    # ==========================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7, BG_LIGHT)
    add_header(s7, "Dual-Score Separation & ATS Quality", "PHASE 7 INNOVATION")

    add_card(s7, Inches(0.8), Inches(1.8), Inches(3.6), Inches(5.0),
             "Dual-Score Separation",
             [
                 "Score A: ATS Quality & Readability (0–100)",
                 "• Evaluates general formatting, structure, and machine readability independent of any specific job.",
                 "Score B: Job-Specific Match (0–100)",
                 "• Evaluates qualification, skills, and experience fit against a target requisition.",
                 "Eliminates confusion between bad formatting and poor qualification fit."
             ], "Architectural Breakthrough")

    add_card(s7, Inches(4.8), Inches(1.8), Inches(3.6), Inches(5.0),
             "7 ATS Quality Categories",
             [
                 "1. Contact & Social Links: 15%",
                 "2. Education Section Formatting: 15%",
                 "3. Experience Section Structure: 20%",
                 "4. Skills Formatting & Taxonomy: 15%",
                 "5. Strong Action Verbs: 15%",
                 "6. Measurable Achievements (% / $): 10%",
                 "7. File Structure & Readability: 10%"
             ], "Quality Matrix")

    add_card(s7, Inches(8.8), Inches(1.8), Inches(3.7), Inches(5.0),
             "Resume Improvement Assistant",
             [
                 "3-Tier prioritized checklist: High, Medium, Low.",
                 "Actionable before/after bullet rewrites.",
                 "Interactive completion toggles saved in database.",
                 "Truthful Writing Guarantee: Suggestions never fabricate unearned degrees, jobs, or metrics."
             ], "Candidate Empowerment")

    # ==========================================================
    # SLIDE 8: CANDIDATE DIGNITY & RECRUITER ANALYTICS
    # ==========================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8, BG_LIGHT)
    add_header(s8, "Candidate Dignity & Recruiter Funnel Analytics", "FEEDBACK & ANALYTICS")

    add_card(s8, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0),
             "'Why Was My Resume Not Shortlisted?'",
             [
                 "Evidence-based gap analysis shows exact missing requirements.",
                 "Side-by-side criteria breakdown (Skills, Years, Qualifications).",
                 "Transparent Disclaimer: 'Decisions reflect the competitive applicant pool for this role.'",
                 "Recruiter Rejection Reason Codes: Qualifications mismatch, experience gap, location constraint, position filled.",
                 "Strict Privacy Protection: Internal recruiter notes and interview rubrics are never exposed to candidates."
             ], "Candidate Respect & Trust")

    add_card(s8, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.0),
             "Recruiter Hiring Funnel & Analytics",
             [
                 "Interactive Chart.js Funnel: Submitted -> Under Review -> Shortlisted -> Interview -> Offered -> Hired.",
                 "Step-by-step conversion rates & drop-off metrics with zero-division safety.",
                 "Time-in-stage and hiring velocity KPIs.",
                 "OWASP Formula Injection Protected CSV Export: Prepends single quotes (') to '=', '+', '-', '@' cells.",
                 "Multi-Tenant Isolation: Recruiters cannot see competitor metrics."
             ], "Data-Driven Hiring")

    # ==========================================================
    # SLIDE 9: QUALITY ASSURANCE & TEST VERIFICATION
    # ==========================================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9, BG_LIGHT)
    add_header(s9, "Comprehensive Test Suite Verification", "QUALITY ASSURANCE")

    add_card(s9, Inches(0.8), Inches(1.8), Inches(3.6), Inches(5.0),
             "Test Suite Overview",
             [
                 "Total Tests: 85 automated tests.",
                 "Pass Rate: 100% (85 passed, 0 failed).",
                 "Execution Time: 27.31 seconds.",
                 "Framework: Pytest 9.0.2 + pytest-flask.",
                 "Database: SQLite in-memory isolation for fast deterministic test runs."
             ], "Pytest Metrics")

    add_card(s9, Inches(4.8), Inches(1.8), Inches(3.6), Inches(5.0),
             "Test Coverage Breakdown",
             [
                 "Auth & Security: 7 tests",
                 "Basic & Health: 5 tests",
                 "Interviews & Pipeline: 14 tests",
                 "Jobs & Applications: 7 tests",
                 "Explainable Matching: 15 tests",
                 "Module Routes: 2 tests",
                 "Phase 7 ATS & Analytics: 14 tests",
                 "Profiles: 3 tests",
                 "RBAC Multi-Tenancy: 6 tests",
                 "Resume Parser & Storage: 12 tests"
             ], "10 Domain Modules")

    add_card(s9, Inches(8.8), Inches(1.8), Inches(3.7), Inches(5.0),
             "Security & Edge Case Hardening",
             [
                 "Multi-tenant isolation verified across all routes.",
                 "IDOR vulnerabilities blocked (HTTP 403 Forbidden).",
                 "Malicious file upload attempts rejected.",
                 "Invalid interview dates & schedule overlaps blocked.",
                 "Zero-division checks on all conversion metrics.",
                 "Spreadsheet formula injection prevented."
             ], "Security Hardened")

    # ==========================================================
    # SLIDE 10: DEMO & HOW TO RUN (Interview Ready)
    # ==========================================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10, NAVY)

    # Header in dark mode
    cat_b = s10.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.5), Inches(0.4))
    p_c = cat_b.text_frame.paragraphs[0]
    p_c.text = "LIVE DEMONSTRATION & REPOSITORY".upper()
    p_c.font.size = Pt(11)
    p_c.font.bold = True
    p_c.font.color.rgb = RGBColor(147, 197, 253)

    t_b = s10.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.5), Inches(0.8))
    p_t = t_b.text_frame.paragraphs[0]
    p_t.text = "How to Run RecruitFlow AI Live for Evaluation"
    p_t.font.size = Pt(28)
    p_t.font.bold = True
    p_t.font.color.rgb = WHITE

    # Left Dark Card
    card_l = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0))
    card_l.fill.solid()
    card_l.fill.fore_color.rgb = RGBColor(30, 41, 59) # Slate 800
    card_l.line.color.rgb = RGBColor(51, 65, 85)

    tb_l = s10.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.2), Inches(4.6))
    tf_l = tb_l.text_frame
    tf_l.word_wrap = True
    p_l1 = tf_l.paragraphs[0]
    p_l1.text = "Quick Launch Commands"
    p_l1.font.size = Pt(18)
    p_l1.font.bold = True
    p_l1.font.color.rgb = RGBColor(56, 189, 248) # Sky blue
    p_l1.space_after = Pt(10)

    cmds = [
        "1. Open Windows PowerShell:",
        "   cd C:\\Users\\sahit\\.gemini\\antigravity\\scratch\\recruitflow-ai",
        "2. Run automated tests (85 passed in ~27s):",
        "   pytest",
        "3. Start application:",
        "   python run.py",
        "4. Open in browser:",
        "   http://127.0.0.1:5000",
        "5. Health API Check:",
        "   http://127.0.0.1:5000/health"
    ]
    for c in cmds:
        p = tf_l.add_paragraph()
        p.text = c
        p.font.size = Pt(12)
        p.font.color.rgb = RGBColor(241, 245, 249)
        p.font.name = "Consolas"
        p.space_after = Pt(4)

    # Right Dark Card
    card_r = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.0))
    card_r.fill.solid()
    card_r.fill.fore_color.rgb = RGBColor(30, 41, 59)
    card_r.line.color.rgb = RGBColor(51, 65, 85)

    tb_r = s10.shapes.add_textbox(Inches(7.0), Inches(2.0), Inches(5.3), Inches(4.6))
    tf_r = tb_r.text_frame
    tf_r.word_wrap = True
    p_r1 = tf_r.paragraphs[0]
    p_r1.text = "Key Demonstration Highlights"
    p_r1.font.size = Pt(18)
    p_r1.font.bold = True
    p_r1.font.color.rgb = RGBColor(56, 189, 248)
    p_r1.space_after = Pt(10)

    highlights = [
        "• Resume Upload & Parsing: Upload a PDF/DOCX resume; observe instant metadata extraction & skills normalization.",
        "• ATS Quality Gauge: View 7-category breakdown and proactive suggestions.",
        "• Side-by-Side Comparison: Compare resume against job description criteria.",
        "• AI Matching & Leaderboard: See match scores calculated deterministically with explainable breakdowns.",
        "• Kanban Pipeline: Move candidate across hiring stages with human-in-the-loop audit logs.",
        "• Recruiter Reports: Inspect Chart.js hiring funnel and export sanitized CSV."
    ]
    for h in highlights:
        p = tf_r.add_paragraph()
        p.text = h
        p.font.size = Pt(12)
        p.font.color.rgb = RGBColor(241, 245, 249)
        p.space_after = Pt(6)

    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

if __name__ == "__main__":
    out_dir = r"C:\Users\sahit\.gemini\antigravity\scratch\recruitflow-ai"
    out_file = os.path.join(out_dir, "RecruitFlow_AI_Presentation.pptx")
    create_deck(out_file)
